"""Builds NVIDIA_Valuation_Model.xlsx with live formulas (assumptions drive DCF)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

GREEN = "76B900"
HDR = PatternFill("solid", fgColor="2E2E2E")
INPUT = PatternFill("solid", fgColor="FFF2CC")   # yellow = editable input
GREENFILL = PatternFill("solid", fgColor="E2EFDA")
BOLD = Font(bold=True)
WHITEBOLD = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, cols, text_map=None):
    for c in cols:
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR
        cell.font = WHITEBOLD
        cell.border = BORDER


# =====================================================================
# Sheet 1: Historicals
# =====================================================================
ws = wb.active
ws.title = "Historicals"
ws["A1"] = "NVIDIA — Historical Financials ($M)"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Source: NVIDIA 10-K filings (FY ends late January). Educational project."
ws["A2"].font = Font(italic=True, size=9)

headers = ["Fiscal Year", "FY2022", "FY2023", "FY2024", "FY2025", "FY2026"]
data = [
    ("Revenue", 26914, 26974, 60922, 130497, 215938),
    ("Cost of revenue", 9439, 11618, 16621, 32639, 62475),
    ("Gross profit", 17475, 15356, 44301, 97858, 153463),
    ("Operating expenses", 7434, 9779, 11329, 16405, 23076),
    ("Operating income", 10041, 5577, 32972, 81453, 130387),
    ("Net income", 9752, 4368, 29760, 72880, 120067),
    ("Diluted EPS ($)", 0.39, 0.17, 1.19, 2.94, 4.90),
    ("Operating cash flow", 9108, 5641, 28090, 64089, 102718),
    ("Capex", 976, 1833, 1069, 3236, 6042),
    ("Free cash flow", 8132, 3808, 27021, 60853, 96676),
    ("Cash & investments", 21208, 13296, 25984, 43210, 62556),
    ("Total debt", 11831, 12031, 11056, 9982, 11412),
    ("Total assets", 44187, 41182, 65728, 111601, 206804),
    ("Total liabilities", 17575, 19081, 22750, 32274, 81196),
    ("Total equity", 26612, 22101, 42978, 79327, 125608),
]
r = 4
for c, h in enumerate(headers, 1):
    ws.cell(row=r, column=c, value=h)
style_header(ws, r, range(1, 7))
r += 1
start = r
for row in data:
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.border = BORDER
        if c == 1:
            cell.font = BOLD
        elif isinstance(val, (int, float)) and row[0] != "Diluted EPS ($)":
            cell.number_format = "#,##0"
        elif row[0] == "Diluted EPS ($)":
            cell.number_format = "0.00"
    r += 1

# derived margin rows (formulas)
ws.cell(row=r, column=1, value="Revenue growth %").font = BOLD
for c in range(3, 7):
    prev = get_column_letter(c - 1)
    cur = get_column_letter(c)
    ws.cell(row=r, column=c, value=f"={cur}{start}/{prev}{start}-1").number_format = "0.0%"
r += 1
ws.cell(row=r, column=1, value="Gross margin %").font = BOLD
for c in range(2, 7):
    col = get_column_letter(c)
    ws.cell(row=r, column=c, value=f"={col}{start+2}/{col}{start}").number_format = "0.0%"
r += 1
ws.cell(row=r, column=1, value="Operating margin %").font = BOLD
for c in range(2, 7):
    col = get_column_letter(c)
    ws.cell(row=r, column=c, value=f"={col}{start+4}/{col}{start}").number_format = "0.0%"
r += 1
ws.cell(row=r, column=1, value="Net margin %").font = BOLD
for c in range(2, 7):
    col = get_column_letter(c)
    ws.cell(row=r, column=c, value=f"={col}{start+5}/{col}{start}").number_format = "0.0%"

ws.column_dimensions["A"].width = 22
for c in range(2, 7):
    ws.column_dimensions[get_column_letter(c)].width = 12

# =====================================================================
# Sheet 2: DCF (live formulas driven by yellow assumption cells)
# =====================================================================
d = wb.create_sheet("DCF")
d["A1"] = "NVIDIA — DCF Valuation"
d["A1"].font = Font(bold=True, size=14)
d["A2"] = "Yellow cells are editable assumptions. Everything else is a formula."
d["A2"].font = Font(italic=True, size=9)

# Assumptions block
d["A4"] = "Assumptions"; d["A4"].font = BOLD
assum = [
    ("Base revenue FY2026 ($M)", 215938, "#,##0"),
    ("Operating margin", 0.60, "0%"),
    ("Tax rate", 0.15, "0%"),
    ("D&A % of revenue", 0.02, "0%"),
    ("Capex % of revenue", 0.03, "0%"),
    ("Incr. NWC % of Δrevenue", 0.03, "0%"),
    ("WACC", 0.10, "0.0%"),
    ("Terminal growth", 0.03, "0.0%"),
    ("Net cash ($M)", 51144, "#,##0"),
    ("Shares outstanding (M)", 24500, "#,##0"),
]
r = 5
for name, val, fmt in assum:
    d.cell(row=r, column=1, value=name)
    cell = d.cell(row=r, column=2, value=val)
    cell.fill = INPUT
    cell.number_format = fmt
    cell.border = BORDER
    r += 1
# name references
BR, OM, TAX, DA, CAPX, NWC, WACC, TG, NC, SH = [f"$B${i}" for i in range(5, 15)]

# growth assumptions row
d["A16"] = "Revenue growth by year"; d["A16"].font = BOLD
growths = [0.45, 0.32, 0.24, 0.18, 0.12]
for i, g in enumerate(growths):
    c = d.cell(row=16, column=2 + i, value=g)
    c.fill = INPUT
    c.number_format = "0%"
    c.border = BORDER

# forecast table
d["A18"] = "Forecast ($M)"; d["A18"].font = BOLD
labels = ["Year", "Revenue", "EBIT", "NOPAT", "+ D&A", "- Capex",
          "- ΔNWC", "FCF", "Discount factor", "PV of FCF"]
for i, lab in enumerate(labels):
    cell = d.cell(row=19, column=1 + i, value=lab)
    cell.fill = HDR
    cell.font = WHITEBOLD
    cell.border = BORDER

for yr in range(1, 6):
    row = 19 + yr
    col_g = get_column_letter(1 + yr)   # growth cell column on row 16
    d.cell(row=row, column=1, value=yr)
    # Revenue: prior * (1+growth). prior year revenue is BR for yr1 else previous B cell
    if yr == 1:
        prev_rev = BR
    else:
        prev_rev = f"B{row-1}"
    d.cell(row=row, column=2, value=f"={prev_rev}*(1+{col_g}16)").number_format = "#,##0"
    d.cell(row=row, column=3, value=f"=B{row}*{OM}").number_format = "#,##0"           # EBIT
    d.cell(row=row, column=4, value=f"=C{row}*(1-{TAX})").number_format = "#,##0"       # NOPAT
    d.cell(row=row, column=5, value=f"=B{row}*{DA}").number_format = "#,##0"            # D&A
    d.cell(row=row, column=6, value=f"=B{row}*{CAPX}").number_format = "#,##0"          # Capex
    if yr == 1:
        d.cell(row=row, column=7, value=f"=(B{row}-{BR})*{NWC}").number_format = "#,##0"
    else:
        d.cell(row=row, column=7, value=f"=(B{row}-B{row-1})*{NWC}").number_format = "#,##0"
    d.cell(row=row, column=8, value=f"=D{row}+E{row}-F{row}-G{row}").number_format = "#,##0"  # FCF
    d.cell(row=row, column=9, value=f"=1/(1+{WACC})^A{row}").number_format = "0.000"
    d.cell(row=row, column=10, value=f"=H{row}*I{row}").number_format = "#,##0"

# outputs
r = 26
outs = [
    ("Sum PV of FCF", "=SUM(J20:J24)"),
    ("Terminal value", f"=H24*(1+{TG})/({WACC}-{TG})"),
    ("PV of terminal value", f"=B27/(1+{WACC})^5"),
    ("Enterprise value", "=B26+B28"),
    ("+ Net cash", f"={NC}"),
    ("Equity value", "=B29+B30"),
    ("÷ Shares (M)", f"={SH}"),
    ("Implied share price ($)", "=B31/B32"),
    ("Current price ($, Sep 9 2026)", 224.15),
    ("Upside / (downside)", "=B33/B34-1"),
]
for name, formula in outs:
    d.cell(row=r, column=1, value=name).font = BOLD
    cell = d.cell(row=r, column=2, value=formula)
    cell.border = BORDER
    if "price" in name.lower() or "value" in name.lower() or "PV" in name or "cash" in name:
        cell.number_format = "#,##0"
    if name.startswith("Implied") or name.startswith("Current"):
        cell.number_format = "$#,##0.00"
        cell.fill = GREENFILL
        cell.font = BOLD
    if name.startswith("Upside"):
        cell.number_format = "0.0%"
        cell.fill = GREENFILL
        cell.font = BOLD
    r += 1

d.column_dimensions["A"].width = 26
for c in range(2, 11):
    d.column_dimensions[get_column_letter(c)].width = 13

# =====================================================================
# Sheet 3: Sensitivity (static values w/ note — Excel data-table optional)
# =====================================================================
s = wb.create_sheet("Sensitivity")
s["A1"] = "DCF Sensitivity — Implied price by WACC × Terminal growth"
s["A1"].font = Font(bold=True, size=13)
s["A2"] = ("Values below are the base-case outputs. To make this fully live, "
           "use Excel's Data Table (What-If Analysis) on the DCF sheet.")
s["A2"].font = Font(italic=True, size=9)

# precomputed base-case grid (matches valuation.py)
tg_headers = ["", "2.0%", "2.5%", "3.0%", "3.5%"]
grid = [
    ["8.0%", 200, 216, 234, 257],
    ["9.0%", 170, 181, 194, 208],
    ["10.0%", 148, 156, 165, 175],
    ["11.0%", 131, 137, 143, 151],
]
for c, h in enumerate(tg_headers, 1):
    cell = s.cell(row=4, column=c, value=h)
    cell.font = WHITEBOLD if c > 1 else BOLD
    if c > 1:
        cell.fill = HDR
for ri, row in enumerate(grid, 5):
    for ci, val in enumerate(row, 1):
        cell = s.cell(row=ri, column=ci, value=val)
        cell.border = BORDER
        if ci == 1:
            cell.fill = HDR
            cell.font = WHITEBOLD
        else:
            cell.number_format = "$#,##0"
            cell.fill = GREENFILL if val >= 224.15 else PatternFill("solid", fgColor="F9D6D6")
s["A11"] = "Green = above current price (~undervalued). Red = below (~overvalued)."
s["A11"].font = Font(italic=True, size=9)
for c in range(1, 6):
    s.column_dimensions[get_column_letter(c)].width = 12

# =====================================================================
# Sheet 4: Comps
# =====================================================================
cp = wb.create_sheet("Comps")
cp["A1"] = "Comparable Companies — Valuation Multiples"
cp["A1"].font = Font(bold=True, size=13)
cp["A2"] = "Source: stockanalysis.com / S&P Global, snapshot as of early Sep 2026."
cp["A2"].font = Font(italic=True, size=9)
comp_head = ["Company", "Ticker", "Trailing P/E", "Forward P/E", "EV/EBITDA",
             "Rev growth", "Op margin"]
comp_rows = [
    ["NVIDIA", "NVDA", 29.1, 19.1, 27.5, "83%", "65%"],
    ["AMD", "AMD", 120.8, 42.6, 79.9, "40%", "20%"],
    ["Broadcom", "AVGO", 61.8, 23.6, 43.1, "45%", "45%"],
    ["Marvell", "MRVL", 79.2, 50.2, 74.7, "45%", "17%"],
]
for c, h in enumerate(comp_head, 1):
    cell = cp.cell(row=4, column=c, value=h)
    cell.fill = HDR; cell.font = WHITEBOLD; cell.border = BORDER
for ri, row in enumerate(comp_rows, 5):
    for ci, val in enumerate(row, 1):
        cell = cp.cell(row=ri, column=ci, value=val)
        cell.border = BORDER
        if isinstance(val, float):
            cell.number_format = "0.0\"x\""
        if ri == 5:
            cell.fill = GREENFILL

# peer-median row and implied valuation (formulas)
cp.cell(row=9, column=1, value="Peer median (excl. NVDA)").font = BOLD
cp.cell(row=9, column=3, value="=MEDIAN(C6:C8)").number_format = "0.0\"x\""
cp.cell(row=9, column=4, value="=MEDIAN(D6:D8)").number_format = "0.0\"x\""
cp.cell(row=9, column=5, value="=MEDIAN(E6:E8)").number_format = "0.0\"x\""

cp.cell(row=11, column=1, value="Implied value from Forward P/E").font = BOLD
cp.cell(row=12, column=1, value="NVIDIA forward EPS ($)")
c = cp.cell(row=12, column=2, value="=224.15/19.1"); c.number_format = "$0.00"
cp.cell(row=13, column=1, value="@ cheapest peer fwd P/E (23.6x)")
c = cp.cell(row=13, column=2, value="=MIN(D6:D8)*B12"); c.number_format = "$#,##0"; c.fill = GREENFILL; c.font = BOLD
cp.cell(row=14, column=1, value="@ peer median fwd P/E")
c = cp.cell(row=14, column=2, value="=D9*B12"); c.number_format = "$#,##0"; c.fill = GREENFILL; c.font = BOLD
cp.cell(row=15, column=1, value="Current price ($)")
c = cp.cell(row=15, column=2, value=224.15); c.number_format = "$#,##0.00"

cp["A17"] = ("Note: NVIDIA has the highest growth and margins yet the LOWEST "
             "multiples of the group. Trailing P/E is distorted for peers (small")
cp["A17"].font = Font(italic=True, size=9)
cp["A18"] = ("current earnings vs. big AI growth expectations), so forward P/E is "
             "the fair comparison. We anchor value near the cheapest peer (~$277).")
cp["A18"].font = Font(italic=True, size=9)
for c in range(1, 8):
    cp.column_dimensions[get_column_letter(c)].width = 15
cp.column_dimensions["A"].width = 28

wb.save("NVIDIA_Valuation_Model.xlsx")
print("Excel model saved")
